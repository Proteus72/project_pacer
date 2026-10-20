# -*- coding: utf-8 -*-
"""计划表解析器：扫描 plans_dirs 下的 .xlsx，按 config.json 的列映射抽取任务行。
原始文件只读，绝不回写。仅支持 .xlsx（.xls/.et 请先另存为 .xlsx）。"""
import datetime as dt
import glob
import os
import re

from openpyxl import load_workbook

from common import resolve
import plan_writer


def _clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def _parse_date(v):
    """把 Excel 单元格值解析成 ISO 日期字符串（YYYY-MM-DD）；失败返回 None"""
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日",
                "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    # Excel 日期序列号（1900 日期系统），粗略范围校验 1954~2119 年
    try:
        n = float(s)
        if 20000 < n < 80000:
            return (dt.date(1899, 12, 30) + dt.timedelta(days=int(n))).isoformat()
    except ValueError:
        pass
    return None


def _norm_status(raw, status_map):
    """状态归一化为 done / doing / todo；返回 (规范状态, 原始文本)"""
    if raw is None:
        return "todo", ""
    s = str(raw).strip()
    if s in status_map:
        return status_map[s], s
    low = s.lower()
    if low in status_map:
        return status_map[low], s
    # 启发式兜底：字面包含关键词
    if "完成" in s or "done" in low or "close" in low:
        return "done", s
    if "进行" in s or "doing" in low or "start" in low or s.endswith("中"):
        return "doing", s
    return "todo", s


def _resolve_columns(header_values, wanted):
    """按表头定位列：先精确匹配，再包含匹配；找不到的列记为 None"""
    headers = [_clean(h) or "" for h in header_values]
    idx = {}
    for logical, name in wanted.items():
        hit = None
        for i, h in enumerate(headers):
            if name and h == name:
                hit = i
                break
        if hit is None:
            for i, h in enumerate(headers):
                if name and name in h:
                    hit = i
                    break
        idx[logical] = hit
    return idx


def _detect_header_row(rows, columns, max_scan=30):
    """自动探测表头行：在前 max_scan 行中按配置列名精确匹配计分，取最高分
    （并列取靠前的一行）；最高分不足 2 时退回第 1 行。
    用于同一模板家族表头行位置不一致的情况（如 AR882 在第 8 行、VASQ 在第 10 行）。"""
    names = [n for n in columns.values() if n]
    best_idx, best_score = 0, 0
    for i, row in enumerate(rows[:max_scan]):
        vals = {_clean(v) or "" for v in row}
        score = sum(1 for n in names if n in vals)
        if score > best_score:
            best_idx, best_score = i, score
    return (best_idx + 1) if best_score >= 2 else 1


def _project_from_cell(rows, cfg):
    """按 project_id_cell（如 'B1'）从工作表元数据区取项目名；
    未配置或取不到时返回 None（由调用方回退到文件名兜底）。"""
    ref = cfg.get("project_id_cell")
    if not ref:
        return None
    m = re.match(r"^([A-Za-z]{1,3})(\d{1,4})$", str(ref).strip())
    if not m:
        return None
    col = 0
    for ch in m.group(1).upper():
        col = col * 26 + (ord(ch) - 64)
    ri, ci = int(m.group(2)) - 1, col - 1
    if 0 <= ri < len(rows) and rows[ri] and ci < len(rows[ri]):
        return _clean(rows[ri][ci])
    return None


def parse_file(path, cfg, fallback_name=None):
    """解析单个 xlsx，返回 (任务 dict 列表, 跳过行数)。
    fallback_name：项目名兜底值（默认取文件名），用于读取回写版本时仍显示原项目名。"""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet_name = cfg.get("sheet_name")
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            if sheet_name:
                print(f"[警告] {os.path.basename(path)} 未找到工作表 '{sheet_name}'，"
                      f"改用第一张 '{wb.worksheets[0].title}'")
            ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    columns = cfg.get("columns", {})
    header_row = int(cfg.get("header_row", 1) or 0)
    if header_row == 0:                       # 0 = 自动探测表头行
        header_row = _detect_header_row(rows, columns)
    if header_row < 1 or header_row > len(rows):
        header_row = 1
    headers = rows[header_row - 1] if rows else []
    col = _resolve_columns(headers, columns)
    status_map = cfg.get("status_map", {})
    require_task = bool(cfg.get("require_task"))
    project_from_cell = _project_from_cell(rows, cfg)
    fallback_project = fallback_name or os.path.splitext(os.path.basename(path))[0]
    # 回写定位用：状态列（1 基）与工作表名
    status_col = (col.get("status") + 1) if col.get("status") is not None else None

    tasks, skipped = [], 0
    for r, row in enumerate(rows[header_row:], start=header_row + 1):
        def g(logical, _row=row, _col=col):
            i = _col.get(logical)
            if i is None or i >= len(_row):
                return None
            return _row[i]

        task = _clean(g("task"))
        subtask = _clean(g("subtask"))
        plan_start = _parse_date(g("plan_start"))
        if require_task and not task:
            continue                          # 阶段汇总行（无任务编号），跳过
        if not (task or subtask):
            continue
        if not plan_start:
            skipped += 1
            continue
        status, raw = _norm_status(g("status"), status_map)
        tasks.append({
            "project": _clean(g("project")) or project_from_cell or fallback_project,
            "task": task or subtask,
            "subtask": subtask,
            "owner": _clean(g("owner")),
            "plan_start": plan_start,
            "plan_end": _parse_date(g("plan_end")),
            "status": status,
            "status_raw": raw,
            "predecessors": _clean(g("predecessors")),
            "source_file": path,
            "source_row": r,
            "sheet": ws.title,          # 回写定位：工作表名
            "status_col": status_col,   # 回写定位：Status 列号（1 基）
            "base_file": fallback_name or os.path.basename(path),
        })
    return tasks, skipped


def parse_all(cfg, quiet=False):
    """扫描所有计划目录，返回 (tasks, 文件数, 跳过行数)。

    每张表优先读「回写的最新版本」（data/plan_versions.json 记录），
    没有回写过才读原表。任务的 base_file 始终是**原表文件名**，
    保证勾选状态 key（文件名|任务|子任务|计划开始）跨版本稳定。"""
    tasks, files, skipped = [], 0, 0
    for d in cfg.get("plans_dirs", []):
        rd = resolve(d)
        for f in sorted(glob.glob(os.path.join(rd, cfg.get("file_glob", "*.xlsx")))):
            if os.path.basename(f).startswith("~$"):
                continue
            src, ver = plan_writer.latest_path(f, cfg)
            ft, sk = parse_file(src, cfg, fallback_name=os.path.basename(f))
            for t in ft:
                t["orig_file"] = f          # 原表完整路径（回写版本链主键）
            tasks.extend(ft)
            skipped += sk
            files += 1
            if not quiet:
                tag = "" if ver == 0 else f" [回写版本 v{ver}]"
                print(f"[解析] {os.path.basename(f)}: {len(ft)} 条任务, "
                      f"跳过 {sk} 行(缺计划开始日期){tag}")
    return tasks, files, skipped
